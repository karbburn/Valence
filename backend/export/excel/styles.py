from __future__ import annotations

"""
Excel styling and convention definitions (openpyxl).

Enforces visual consistency across all workbook tabs:
- Blue font = user input / editable assumption
- Black font = formula / calculated cell
- Green font = cross-sheet link
- Red font = QA failure / alert flag
- Consistent number formats (negatives in parentheses, INR Crores)
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Font definitions
FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="1F4E79")
FONT_SUBTITLE = Font(name="Calibri", size=11, italic=True, color="595959")
FONT_SECTION = Font(name="Calibri", size=12, bold=True, color="1F4E79")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_SUBHEADER = Font(name="Calibri", size=11, bold=True, color="000000")

FONT_INPUT = Font(name="Calibri", size=11, color="1F4E79")         # Blue font = input
FONT_FORMULA = Font(name="Calibri", size=11, color="000000")       # Black font = formula
FONT_LINK = Font(name="Calibri", size=11, color="375623")          # Green font = cross-sheet link
FONT_HYPERLINK = Font(name="Calibri", size=11, underline="single", color="0563C1") # Hyperlink blue
FONT_ALERT = Font(name="Calibri", size=11, bold=True, color="C00000") # Red font = alert/failure
FONT_PASS = Font(name="Calibri", size=11, bold=True, color="375623")  # Green font = pass
FONT_TOTAL = Font(name="Calibri", size=11, bold=True, color="000000") # Bold total line

# Fills
FILL_HEADER = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
FILL_CARD = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
FILL_PASS = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
FILL_FAIL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

# Borders
SIDE_THIN = Side(style="thin", color="D9D9D9")
SIDE_DOUBLE = Side(style="double", color="000000")
SIDE_TOTAL_TOP = Side(style="thin", color="000000")

BORDER_BOX = Border(left=SIDE_THIN, right=SIDE_THIN, top=SIDE_THIN, bottom=SIDE_THIN)
BORDER_HEADER = Border(left=SIDE_THIN, right=SIDE_THIN, top=SIDE_THIN, bottom=Side(style="medium", color="1F4E79"))
BORDER_TOTAL = Border(top=SIDE_TOTAL_TOP, bottom=SIDE_DOUBLE)

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

# Number formats
FMT_AMOUNT = "#,##0.0;(#,##0.0);\"-\""
FMT_INT = "#,##0;(#,##0);\"-\""
FMT_CURRENCY_INT = "#,##0;(#,##0);\"-\""
FMT_PERCENT = "0.0%"
FMT_PERCENT_PRECISION = "0.00%"
FMT_MULTIPLE = "0.0\"x\""
FMT_DAYS = "0.0\" days\""
FMT_PRICE = "#,##0.00"
