"""Stock-based compensation flows from ingestion through forecast into FCFF."""

from datetime import date

import pytest
from openpyxl import Workbook

from backend.export.excel.render_val import render_dcf_tab
from backend.forecast.engine import run_forecast
from backend.models.spec.metadata import ModelMetadata
from backend.models.spec.historicals import Historicals
from backend.models.spec.model_specification import ModelSpecification
from backend.models.spec.valuation import (
    DCFBridge,
    FCFFPeriod,
    TerminalValue,
    ValuationOutput,
    WACCBreakdown,
)
from backend.models.statements.historical_model import build_historical_model
from backend.normalization.taxonomy.models import CanonicalDatapoint
from backend.valuation.dcf import compute_fcff_periods


def _dp(key: str, period: str, value: float) -> CanonicalDatapoint:
    return CanonicalDatapoint(
        id=f"{key}-{period}",
        company_id="sbc_us",
        canonical_key=key,
        metric_raw=key,
        period_label=period,
        period_end_date=date(2026, 3, 31),
        value=float(value),
        currency="USD",
        units="millions",
        status="reported",
        source_datapoint_ids=["fixture"],
    )


@pytest.fixture()
def historical_with_sbc():
    datapoints = []
    for idx, period in enumerate(["FY25", "FY26"]):
        revenue = 100_000.0 * (1.1 ** idx)
        datapoints.append(_dp("canonical.is.revenue", period, revenue))
        datapoints.append(_dp("canonical.cf.stock_compensation", period, 1_000.0 * (1.1 ** idx)))
    return build_historical_model(datapoints, target_periods=["FY25", "FY26"])


def test_engine_projects_sbc_at_historical_pct_of_revenue(historical_with_sbc):
    forecast = run_forecast([], historical_with_sbc, "base")
    values = {(li.canonical_key, li.period_label): li.value for li in forecast.line_items}
    # FY26 SBC / revenue = 1%; FY27 revenue = 110000 * 1.10 -> SBC 1210
    assert values[("canonical.cf.stock_compensation", "FY27")] == pytest.approx(1210.0)


def test_fcff_deducts_stock_compensation(historical_with_sbc):
    forecast = run_forecast([], historical_with_sbc, "base")
    periods = compute_fcff_periods(forecast, wacc_pct=12.0, scenario="base")
    for p in periods:
        expected = (
            (p.nopat or 0.0)
            + (p.da or 0.0)
            - (p.capex or 0.0)
            - (p.delta_working_capital or 0.0)
            - (p.stock_compensation or 0.0)
        )
        assert p.fcff == pytest.approx(expected, abs=0.05)
    assert all((p.stock_compensation or 0.0) > 0 for p in periods)


def test_dcf_without_sbc_data_leaves_fcff_untouched():
    class _Empty:
        def get_value(self, *_args, **_kwargs):
            return None

    forecast = _Empty()
    periods = compute_fcff_periods(forecast, wacc_pct=12.0, scenario="base")
    assert len(periods) == 5
    assert all(p.stock_compensation == 0.0 for p in periods)


def _spec_with_sbc(sbc_value: float) -> ModelSpecification:
    fcff_periods = [
        FCFFPeriod(
            period=p,
            ebit=100.0,
            tax_rate=21.0,
            nopat=79.0,
            da=10.0,
            capex=15.0,
            delta_working_capital=2.0,
            stock_compensation=sbc_value,
            fcff=79.0 + 10.0 - 15.0 - 2.0 - sbc_value,
            discount_factor=0.9,
            pv_fcff=(79.0 + 10.0 - 15.0 - 2.0 - sbc_value) * 0.9,
        )
        for p in ["FY27", "FY28"]
    ]
    return ModelSpecification(
        metadata=ModelMetadata(
            company_id="sbc_us",
            ticker="SBC",
            name="SBC Testco",
            market="us",
            currency="USD",
            units="millions",
            fiscal_year_end="Dec 31",
            shares_outstanding=100.0,
        ),
        historicals=Historicals(periods=[], line_items=[]),
        valuation=[
            ValuationOutput(
                scenario="base",
                wacc=WACCBreakdown(wacc=12.0),
                fcff_by_period=fcff_periods,
                terminal_value=TerminalValue(terminal_growth_rate=4.0, exit_multiple=20.0),
                dcf_bridge=DCFBridge(shares_outstanding=100.0),
            )
        ],
    )


def test_excel_dcf_tab_renders_memo_row_and_extended_fcff_formula():
    spec = _spec_with_sbc(sbc_value=8.0)
    ws = render_dcf_tab(Workbook(), spec)

    label = ws.cell(row=15, column=2).value
    cached = ws.cell(row=15, column=3).value
    formula = ws.cell(row=12, column=3).value

    assert label == "Less: Stock-Based Comp (Memo)"
    assert cached == 8.0
    assert formula == "=C8+C9-C10-C11-C15"


def test_excel_dcf_tab_defaults_missing_sbc_to_zero():
    spec = _spec_with_sbc(sbc_value=0.0)
    ws = render_dcf_tab(Workbook(), spec)
    assert ws.cell(row=15, column=3).value == 0.0
