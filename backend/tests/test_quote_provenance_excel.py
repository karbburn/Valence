"""Excel export must carry market-quote provenance (date + source), not a bare price.

An exported model is read days later by someone who has no way to know whether
the "Market Benchmark Price" was a live close or a stale fallback. These tests
pin the provenance stamp in 34_Reverse_DCF and 02_Executive_Summary.
"""

from openpyxl import Workbook

from backend.export.excel.render_front import render_executive_summary
from backend.export.excel.render_val import render_reverse_dcf_tab
from backend.models.spec.historicals import Historicals
from backend.models.spec.metadata import ModelMetadata
from backend.models.spec.model_specification import ModelSpecification
from backend.models.spec.valuation import (
    DCFBridge,
    ReverseDCF,
    ValuationOutput,
    WACCBreakdown,
)


def _spec(price=225.07, quote_date="2026-09-25", source="yfinance_history") -> ModelSpecification:
    spec = ModelSpecification(
        metadata=ModelMetadata(
            company_id="prov_us",
            ticker="PRV",
            name="Provenance Testco",
            market="us",
            currency="USD",
            units="millions",
            fiscal_year_end="December 31",
            shares_outstanding=1000.0,
        ),
        historicals=Historicals(periods=["FY26"], line_items=[]),
    )
    spec.valuation = [
        ValuationOutput(
            scenario=sc,
            wacc=WACCBreakdown(wacc=9.0),
            dcf_bridge=DCFBridge(
                enterprise_value=100_000.0,
                equity_value=90_000.0,
                shares_outstanding=1000.0,
                implied_share_price=90.0,
                less_net_debt=10_000.0,
                sum_pv_fcff=40_000.0,
            ),
            reverse_dcf=ReverseDCF(
                market_price=price,
                market_price_date=quote_date,
                market_price_source=source,
                implied_terminal_growth=3.0,
            ),
        )
        for sc in ("base", "bull", "bear")
    ]
    return spec


def test_reverse_dcf_sheet_shows_quote_date_and_source():
    ws = render_reverse_dcf_tab(Workbook(), _spec())
    assert ws["C6"].value == 225.07

    prov_rows = [
        r for r in range(1, ws.max_row + 1)
        if str(ws.cell(row=r, column=2).value) == "Quote as of / source"
    ]
    assert prov_rows, "34_Reverse_DCF must carry a quote provenance row"
    stamp = str(ws.cell(row=prov_rows[0], column=3).value)
    assert "2026-09-25" in stamp and "yfinance_history" in stamp
    assert "STALE" not in stamp.upper()


def test_reverse_dcf_flags_a_stale_quote():
    ws = render_reverse_dcf_tab(
        Workbook(), _spec(price=1080.0, quote_date="2026-08-14", source="stale_cache:yfinance")
    )
    stamps = [
        str(ws.cell(row=r, column=3).value)
        for r in range(1, ws.max_row + 1)
        if str(ws.cell(row=r, column=2).value) == "Quote as of / source"
    ]
    assert stamps and "STALE" in stamps[0].upper()
    assert "2026-08-14" in stamps[0]


def test_executive_summary_kpi_label_includes_provenance():
    ws = render_executive_summary(Workbook(), _spec())
    assert ws["B6"].value == 225.07
    label = str(ws["B5"].value)
    assert "2026-09-25" in label and "yfinance_history" in label
