"""Contract between rendered Excel layouts and every formula that references them.

These tests fail loudly when a renderer shifts a row/column that live formulas
in other tabs depend on — the class of silent breakage this suite exists to stop.
"""

import re

from openpyxl import Workbook

from backend.export.excel.links import AUTHOR_URL, VALENCE_URL
from backend.export.excel.render_fcst import render_debt_schedule
from backend.export.excel.render_front import render_cover
from backend.export.excel.render_hist import (
    render_historical_balance_sheet,
    render_historical_cash_flow,
)
from backend.export.excel.render_qa import render_model_checks_tab
from backend.export.excel.render_val import render_dcf_tab, render_sensitivity_tab
from backend.models.spec.historicals import HistoricalLineItem, Historicals
from backend.models.spec.metadata import ModelMetadata
from backend.models.spec.model_specification import ModelSpecification
from backend.models.spec.qa import ModelCheckResult, QAResults

PERIODS = ["FY24", "FY25", "FY26"]


def _hist_line(key: str, value: float) -> HistoricalLineItem:
    return HistoricalLineItem(
        canonical_key=key,
        period_label=PERIODS[-1],
        period_end_date=__import__("datetime").date(2026, 3, 31),
        value=value,
        currency="INR",
        units="crores",
        status="reported",
        source_datapoint_ids=["fixture"],
        values_by_period={p: value for p in PERIODS},
    )


def _spec_with_historicals(bs_keys, cf_keys) -> ModelSpecification:
    lines = [_hist_line(k, 100.0) for k in bs_keys + cf_keys]
    return ModelSpecification(
        metadata=ModelMetadata(
            company_id="contract_in",
            ticker="CTRT",
            name="Contract Testco",
            market="india",
            currency="INR",
            units="crores",
            fiscal_year_end="March 31",
            shares_outstanding=10.0,
        ),
        historicals=Historicals(periods=PERIODS, line_items=lines),
    )


BS_KEYS = [
    "canonical.bs.ppe",
    "canonical.bs.cwip",
    "canonical.bs.goodwill",
    "canonical.bs.intangible_assets",
    "canonical.bs.non_current_investments",
    "canonical.bs.deferred_tax_assets",
    "canonical.bs.total_non_current_assets",
    "canonical.bs.trade_receivables",
    "canonical.bs.unbilled_revenue",
    "canonical.bs.inventory",
    "canonical.bs.cash_and_bank",
    "canonical.bs.current_investments",
    "canonical.bs.total_current_assets",
    "canonical.bs.total_assets",
    "canonical.bs.equity_capital",
    "canonical.bs.total_equity",
    "canonical.bs.trade_payables",
    "canonical.bs.total_liabilities_and_equity",
]

CF_KEYS = [
    "canonical.cf.operating_activities",
    "canonical.cf.taxes_paid",
    "canonical.cf.investing_activities",
    "canonical.cf.business_acquisitions",
    "canonical.cf.financing_activities",
    "canonical.cf.dividends_paid",
    "canonical.cf.other_adjustments",
    "canonical.cf.net_change_in_cash",
]


def _row_of(ws, label: str) -> int:
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=2).value) == label:
            return r
    raise AssertionError(f"Label not found on sheet: {label}")


def test_cover_links_match_current_public_urls():
    ws = render_cover(Workbook(), _spec_with_historicals([], []))

    assert ws["B20"].value == "By Sourabh"
    assert ws["B20"].hyperlink is not None
    assert ws["B20"].hyperlink.target == AUTHOR_URL
    # The wordmark is the single place the cover navigates to the platform. The
    # strapline beneath it is plain text, so the cover offers one target rather
    # than two adjacent cells linking to the same place.
    assert ws["B2"].value == "V A L E N C E"
    assert ws["B2"].hyperlink is not None
    assert ws["B2"].hyperlink.target == VALENCE_URL
    assert ws["B3"].hyperlink is None, (
        "the strapline under the wordmark is also a link, so the cover presents "
        "two links to the same destination"
    )


def test_balance_sheet_layout_matches_formula_references():
    ws = render_historical_balance_sheet(Workbook(), _spec_with_historicals(BS_KEYS, []))
    # Pinned as an ordering, not as row numbers. These rows are referenced by the
    # audit formulas on 52_Model_Checks, but those formulas resolve their rows by
    # label, so what matters here is that the statement reads in the right order
    # and that every line the checks need is present.
    order = [
        "Property, Plant & Equipment",
        "Goodwill",
        "Intangible Assets",
        "Non-Current Investments",
        "Deferred Tax Assets",
        "Other Non-Current Assets",
        "Total Non-Current Assets",
        "Trade Receivables",
        "Unbilled Revenue",
        "Inventory",
        "Cash & Cash Equivalents",
        "Current Investments",
        "Total Current Assets",
        "TOTAL ASSETS",
        "Total Equity / Net Worth",
        "Trade Payables",
        "TOTAL LIABILITIES & EQUITY",
    ]
    rows = [_row_of(ws, label) for label in order]
    assert rows == sorted(rows), "the balance sheet is not in filed order"
    assert len(set(rows)) == len(rows), "two lines share a row"


def test_dcf_working_capital_reads_the_working_capital_lines_by_label():
    """Year-one working capital must open from the four lines it names.

    These were rows 13, 14, 15 and 22, describing the balance sheet as it stood
    when they were written. Adding a line to that statement moved every row beneath
    it, so the formula read the non-current subtotal as trade receivables, left
    inventory out of the working capital base entirely, and subtracted total
    equity in place of trade payables. It still returned a number, so the workbook
    showed the engine's cached figure on open and a different one after
    recalculation. This is the formula that feeds the first forecast year's change
    in working capital, and through it free cash flow and the enterprise value.
    """
    wb = Workbook()
    spec = _spec_with_historicals(BS_KEYS, CF_KEYS)
    bs = render_historical_balance_sheet(wb, spec)
    render_dcf_tab(wb, spec)

    dcf = wb["31_DCF"]
    text = _all_formula_text(dcf)
    for label in ("Trade Receivables", "Unbilled Revenue", "Inventory", "Trade Payables"):
        row = _row_of(bs, label)
        assert re.search(rf"11_Balance_Sheet'!\w+{row}\b", text), (
            f"the year-one working capital formula does not reference {label} at "
            f"its actual row {row}"
        )

    # And nothing else on the balance sheet may be pulled into the base.
    referenced = {int(m) for m in re.findall(r"11_Balance_Sheet'!\w{1,2}?(\d+)\b", text)}
    expected = {
        _row_of(bs, "Trade Receivables"),
        _row_of(bs, "Unbilled Revenue"),
        _row_of(bs, "Inventory"),
        _row_of(bs, "Trade Payables"),
    }
    assert referenced == expected, (
        f"the formula references balance sheet rows {sorted(referenced)} against the "
        f"four working capital lines {sorted(expected)}"
    )


def test_audit_formulas_follow_the_balance_sheet_lines_they_name():
    """The reconciliation checks must reference the rows they claim to.

    They used to name rows by number, so adding a line to the balance sheet moved
    the totals underneath them and the checks compared unrelated cells while still
    reporting a verdict. This pins the formulas to the labelled lines.
    """
    wb = Workbook()
    spec = _spec_with_historicals(BS_KEYS, CF_KEYS)
    spec.qa = QAResults(
        checks=[
            ModelCheckResult(check_name="balance_sheet_balances", category="accounting", passed=True),
            ModelCheckResult(check_name="cash_flow_reconciles", category="accounting", passed=True),
        ]
    )
    bs = render_historical_balance_sheet(wb, spec)
    render_historical_cash_flow(wb, spec)
    qa = render_model_checks_tab(wb, spec)

    text = _all_formula_text(qa)
    # The cash check reads this year's balance and last year's, so any column is
    # acceptable; what matters is that the row is the one the line occupies.

    for label in ("TOTAL ASSETS", "TOTAL LIABILITIES & EQUITY", "Cash & Cash Equivalents"):
        row = _row_of(bs, label)
        assert re.search(rf"11_Balance_Sheet'!\w+{row}\b", text), (
            f"no audit formula references {label} at its actual row {row}"
        )


def _all_formula_text(ws) -> str:
    parts = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                parts.append(cell.value)
    return " ".join(parts)


def test_cash_flow_layout_keeps_net_change_on_row_12():
    ws = render_historical_cash_flow(Workbook(), _spec_with_historicals([], CF_KEYS))
    assert _row_of(ws, "NET CHANGE IN CASH & CASH EQUIVALENTS") == 12
    assert _row_of(ws, "Memo: Stock-Based Compensation") == 13


def test_debt_schedule_rows_support_internal_reconciliation_formula():
    ws = render_debt_schedule(Workbook(), _spec_with_historicals([], []))
    assert _row_of(ws, "Opening Debt Balance (INR Cr)") == 6
    assert _row_of(ws, "Debt Drawdowns (INR Cr)") == 7
    assert _row_of(ws, "Scheduled Repayments (INR Cr)") == 8
    assert _row_of(ws, "Optional Repayments (INR Cr)") == 9
    closing_formula = ws.cell(row=10, column=3).value
    assert closing_formula == "=C6+C7-C8-C9"


def test_sensitivity_formulas_rediscount_the_full_fcff_stream():
    ws = render_sensitivity_tab(Workbook(), _spec_with_historicals([], []))

    growth_cell = ws.cell(row=6, column=3).value  # first WACC row x first growth col
    assert isinstance(growth_cell, str) and growth_cell.startswith("=IF(")
    for t in ("0.5", "1.5", "2.5", "3.5", "4.5"):
        assert f"(1+$B6)^{t}" in growth_cell, f"missing trial-WACC exponent {t}"
    assert "'31_DCF'!C12" in growth_cell and "'31_DCF'!G12/(1+$B6)^4.5" in growth_cell

    multiple_cell = ws.cell(row=14, column=3).value
    assert isinstance(multiple_cell, str) and multiple_cell.startswith("=(")
    assert "(1+$B14)^0.5" in multiple_cell and "'20_Operating_Model'!G9*C$13" in multiple_cell


def test_model_check_formulas_point_at_rendered_cells():
    spec = _spec_with_historicals(BS_KEYS, CF_KEYS)
    spec.qa = QAResults(
        checks=[
            ModelCheckResult(check_name="balance_sheet_balances", category="accounting", passed=True),
            ModelCheckResult(check_name="cash_flow_reconciles", category="accounting", passed=True),
            ModelCheckResult(check_name="debt_schedule_reconciles", category="accounting", passed=True),
        ]
    )
    wb = Workbook()
    bs = render_historical_balance_sheet(wb, spec)
    render_historical_cash_flow(wb, spec)
    ws = render_model_checks_tab(wb, spec)

    formulas = {}
    for r in range(1, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        val = ws.cell(row=r, column=4).value or ws.cell(row=r, column=5).value
        if isinstance(val, str) and val.startswith("=IF"):
            formulas[str(name)] = val

    # Referenced at the rows the balance sheet actually puts those lines on, which
    # is not a fixed number, and is exactly why the formulas resolve them by label.
    assets_row = _row_of(bs, "TOTAL ASSETS")
    le_row = _row_of(bs, "TOTAL LIABILITIES & EQUITY")
    cash_row = _row_of(bs, "Cash & Cash Equivalents")

    bs_formula = formulas["balance_sheet_balances"]
    assert f"'11_Balance_Sheet'!E{assets_row}" in bs_formula
    assert f"'11_Balance_Sheet'!E{le_row}" in bs_formula

    cf_formula = formulas["cash_flow_reconciles"]
    assert "'12_Cash_Flow'!E12" in cf_formula
    assert f"'11_Balance_Sheet'!E{cash_row}" in cf_formula

    # The tie-out is only meaningful with a prior period to subtract, and a single
    # period cannot produce one. Asserted here because the row the prior column
    # points at has to be the same cash row, not a row that shifted with the layout.
    two = _spec_with_historicals(BS_KEYS, CF_KEYS)
    two.historicals.periods = ["FY25", "FY26"]
    wb2 = Workbook()
    bs2 = render_historical_balance_sheet(wb2, two)
    render_historical_cash_flow(wb2, two)
    two.qa = QAResults(
        checks=[
            ModelCheckResult(check_name="cash_flow_reconciles", category="accounting", passed=True),
        ]
    )
    qa2 = render_model_checks_tab(wb2, two)
    cash_row2 = _row_of(bs2, "Cash & Cash Equivalents")
    text2 = _all_formula_text(qa2)
    assert re.search(rf"11_Balance_Sheet'!D{cash_row2}\b", text2), (
        "the cash tie-out does not subtract the prior year's balance at the row the "
        "cash line actually occupies"
    )
    # No reference may point anywhere else on the balance sheet. The prior-column
    # reference was left hardcoded while the closing one was resolved, so it
    # silently compared this year's cash with last year's inventory.
    referenced = {int(m) for m in re.findall(r"11_Balance_Sheet'!\w{1,2}?(\d+)\b", text2)}
    assert referenced == {cash_row2}, (
        f"the cash tie-out references balance sheet rows {sorted(referenced)}, but "
        f"only the cash line at row {cash_row2} belongs in it"
    )

    debt_formula = formulas["debt_schedule_reconciles"]
    for col in "CDEFG":
        assert f"'25_Debt_Schedule'!{col}10" in debt_formula
