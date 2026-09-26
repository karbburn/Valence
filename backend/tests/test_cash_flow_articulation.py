"""The cash flow statement must add up, and say so when it does not.

Two defects this pins.

Net change in cash was read as a fourth reported figure rather than derived as
the sum of the three sections — which is its definition. The feeds report it as
zero, so the cash flow statement published a net change in cash of 0.0 for every
company in every period while the balance sheet's own cash line moved: at one
large-cap by 1,309 and 2,016, and at another by 1,700 and 2,300. The line a
reader checks to see whether the statement adds up was the one line guaranteed
not to.

The QA tab cached the engine's verdict unconditionally, so `debt_schedule_reconciles`
showed PASS on open and FAIL after a recalculation — its formula compared the
CLOSING debt balance against the OPTIONAL REPAYMENTS row, which for a company
that carries debt and repays nothing optionally is 8,468 against 0. A cached value
that disagrees with the formula beside it makes the workbook contradict itself
depending on when it was opened.
"""

from __future__ import annotations

import pytest

from backend.export.excel.render_qa import render_model_checks_tab
from backend.models.statements.balance_sheet import BalanceSheet
from backend.models.statements.cash_flow import (
    CashFlowLineItem,
    CashFlowStatement,
    derive_net_change_in_cash,
)


def _cf(periods, operating, investing, financing, reported_net=None):
    def _values(seq):
        # A None means "the feed did not report this period", which is different
        # from a reported zero and must not become a value.
        return {
            p: v for p, v in zip(periods, seq) if v is not None
        }

    items = [
        CashFlowLineItem(
            canonical_key=key,
            display_label=label,
            category=cat,
            values_by_period=_values(values),
        )
        for key, label, cat, values in (
            ("canonical.cf.operating_activities", "CFO", "operating", operating),
            ("canonical.cf.investing_activities", "CFI", "investing", investing),
            ("canonical.cf.dividends_paid", "Dividends", "financing", [0.0] * len(periods)),
            ("canonical.cf.financing_activities", "CFF", "financing", financing),
        )
    ]
    if reported_net is not None:
        items.append(
            CashFlowLineItem(
                canonical_key="canonical.cf.net_change_in_cash",
                display_label="Net Change in Cash",
                category="summary",
                values_by_period=_values(reported_net),
            )
        )
    return CashFlowStatement(company_id="t", periods=periods, line_items=items)


PERIODS = ["FY24", "FY25", "FY26"]


# --------------------------------------------------------------------------- #
# Net change in cash is the sum of the three sections
# --------------------------------------------------------------------------- #

def test_a_reported_zero_is_replaced_by_the_sum():
    """The feeds report zero; the balance sheet's cash line does not stand still."""
    cf = _cf(PERIODS, [28_090.0, 64_089.0, 102_718.0],
             [-10_566.0, -20_421.0, -52_228.0],
             [-13_633.0, -42_359.0, -48_474.0],
             reported_net=[0.0, 0.0, 0.0])
    derive_net_change_in_cash(cf)

    assert cf.get_value("canonical.cf.net_change_in_cash", "FY25") == pytest.approx(1_309.0)
    assert cf.get_value("canonical.cf.net_change_in_cash", "FY26") == pytest.approx(2_016.0)
    assert cf.get_value("canonical.cf.net_change_in_cash", "FY24") == pytest.approx(3_891.0)


def test_the_line_is_created_when_the_source_never_reported_one():
    """A missing line is worse than a wrong one: the reader cannot check at all."""
    cf = _cf(PERIODS, [100.0, 120.0, 140.0], [-10.0, -12.0, -14.0], [-60.0, -58.0, -56.0])
    assert cf.get_value("canonical.cf.net_change_in_cash", "FY25") is None

    derive_net_change_in_cash(cf)

    assert cf.get_value("canonical.cf.net_change_in_cash", "FY25") == pytest.approx(50.0)


def test_the_line_is_the_total_of_the_statement():
    """It must be the sum, so a reader adding the three sections gets the same number."""
    cf = _cf(PERIODS, [100.0, 120.0, 140.0], [-10.0, -12.0, -14.0], [-60.0, -58.0, -56.0])
    derive_net_change_in_cash(cf)

    for p in PERIODS:
        sections = sum(
            cf.get_value(k, p)
            for k in (
                "canonical.cf.operating_activities",
                "canonical.cf.investing_activities",
                "canonical.cf.financing_activities",
            )
        )
        assert cf.get_value("canonical.cf.net_change_in_cash", p) == pytest.approx(sections)


def test_a_partial_statement_is_not_given_an_invented_total():
    """Three sections minus one is not a net change, and must not be published as one."""
    cf = _cf(PERIODS, [100.0, 120.0, 140.0], [-10.0, -12.0, -14.0], [None] * 3)
    derive_net_change_in_cash(cf)

    assert cf.get_value("canonical.cf.net_change_in_cash", "FY25") is None


def test_a_negative_net_change_is_published_as_negative():
    """A company that consumed cash consumed it. A truthiness test would zero it."""
    cf = _cf(PERIODS, [50.0, 60.0, 70.0], [-10.0, -12.0, -14.0], [-100.0, -98.0, -96.0],
             reported_net=[0.0, 0.0, 0.0])
    derive_net_change_in_cash(cf)

    assert cf.get_value("canonical.cf.net_change_in_cash", "FY25") == pytest.approx(-50.0)
    assert cf.get_value("canonical.cf.net_change_in_cash", "FY25") < 0


# --------------------------------------------------------------------------- #
# A cached verdict must not disagree with the formula beside it
# --------------------------------------------------------------------------- #

# The shared `exported` fixture does not run the QA pass, so its workbook has an
# empty checks tab. These tests are about the checks, so they build a spec that
# has them.
from backend.tests.test_excel_recalculation_parity import _model  # noqa: E402


@pytest.fixture(scope="module")
def checked_workbook(tmp_path_factory):
    """A workbook whose QA tab is populated."""
    from openpyxl import load_workbook

    from backend.export.excel.exporter import export_model_to_excel
    from backend.validation.pipeline import run_qa
    from backend.valuation.pipeline import run_valuation

    spec = run_qa(run_valuation(_model()))
    path = tmp_path_factory.mktemp("qa") / "checked.xlsx"
    export_model_to_excel(spec, path)
    return (
        load_workbook(path, data_only=True),
        load_workbook(path, data_only=False),
    )


def _qa_rows(wb):
    """(name, verdict-or-formula) for every check on the QA tab.

    Scans the whole sheet rather than a fixed row range: the tab's layout is not
    a contract, and a test that breaks when a title moves is testing the layout
    instead of the checks.
    """
    ws = wb["52_Model_Checks"]
    rows = []
    for r in range(1, (ws.max_row or 0) + 1):
        name = ws.cell(row=r, column=2).value
        if not name or str(name).strip() in ("Check Name", "Line Item"):
            continue
        rows.append((str(name), ws.cell(row=r, column=4).value))
    return rows


def test_the_debt_check_compares_the_right_rows(checked_workbook):
    """Closing debt is opening plus drawdowns less repayments.

    It compared the closing balance against the OPTIONAL REPAYMENTS row, so a
    company carrying debt and repaying nothing optionally was checked as, say,
    8,468 against 0 — and returned FAIL in every workbook while the tab said PASS.
    """
    import re

    _, wbf = checked_workbook
    formula = dict(_qa_rows(wbf))["debt_schedule_reconciles"]

    assert isinstance(formula, str) and formula.startswith("="), formula

    # Every term of the identity is referenced, for every forecast column.
    for col in "CDEFG":
        for row in ("6", "7", "8", "9", "10"):
            assert f"!{col}{row}" in formula, (
                f"the debt check does not reference {col}{row}, so it is not testing "
                f"the closing-debt identity: {formula}"
            )

    # And it must not be the bare comparison it used to be.
    assert not re.search(r"!\w10-N\('25_Debt_Schedule'!\w9\)\)", formula), (
        f"the check still compares closing debt to optional repayments: {formula}"
    )


def test_a_check_with_no_input_reports_a_failure_rather_than_a_pass(checked_workbook):
    """Absence is not a pass.

    The fixture carries no debt schedule, so the debt check has nothing to
    verify. Reporting PASS for a check that could not run is the same class of
    defect as reporting PASS for one that ran and failed — and a reader cannot
    tell the two apart from the verdict alone.
    """
    wbv, wbf = checked_workbook
    ws = wbv["52_Model_Checks"]
    verdicts = dict(_qa_rows(wbf))

    detail = None
    for r in range(1, (ws.max_row or 0) + 1):
        if str(ws.cell(row=r, column=2).value or "") == "debt_schedule_reconciles":
            detail = str(ws.cell(row=r, column=5).value or "")
            break

    assert detail and "no debt schedule" in detail.lower(), (
        f"expected the check to say it had no debt schedule to verify, got {detail!r}"
    )
    assert verdicts.get("debt_schedule_reconciles") != "PASS", (
        "a check with nothing to verify must not report PASS"
    )


def test_no_cached_verdict_contradicts_its_own_detail(checked_workbook):
    """A row must not say PASS while its own explanation says otherwise.

    The cash flow statement does not articulate for several companies and the
    engine's detail says so; the tab used to publish PASS regardless, which is a
    workbook that disagrees with itself.
    """
    wbv, wbf = checked_workbook
    ws = wbv["52_Model_Checks"]
    contradictions = []
    for r in range(1, (ws.max_row or 0) + 1):
        name = ws.cell(row=r, column=2).value
        verdict = ws.cell(row=r, column=4).value
        detail = str(ws.cell(row=r, column=5).value or "").lower()
        if not name or verdict != "PASS":
            continue
        if any(p in detail for p in ("does not reconcile", "does not balance", "not reconciled")):
            contradictions.append(f"{name}: PASS but detail says {detail!r}")

    assert not contradictions, contradictions
