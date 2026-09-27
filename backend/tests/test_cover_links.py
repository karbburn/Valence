"""The cover must carry a working link to the platform.

The wordmark on the cover is the element a reader reaches for. It was not
hyperlinked while the strapline beneath it was, so the cover looked like it had
no destination at all — and nothing checked, so a renderer change could drop it
without a word.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from backend.export.excel.links import VALENCE_URL
from backend.export.excel.exporter import export_model_to_excel
from backend.models.spec.model_specification import ModelSpecification
from backend.valuation.pipeline import run_valuation


@pytest.fixture(scope="module")
def cover(tmp_path_factory):
    from backend.tests.test_excel_recalculation_parity import _model

    spec = run_valuation(_model())
    path = tmp_path_factory.mktemp("cover") / "cover.xlsx"
    export_model_to_excel(spec, path)
    return load_workbook(path)["00_Cover"]


def test_the_wordmark_on_the_cover_links_to_the_platform(cover):
    """B2 is the brand; it must navigate."""
    cell = cover["B2"]

    assert str(cell.value).replace(" ", "") == "VALENCE", cell.value
    assert cell.hyperlink is not None, (
        "the wordmark on the cover carries no hyperlink, so the cover's most "
        "reachable element does nothing when clicked"
    )
    assert cell.hyperlink.target == VALENCE_URL, (
        f"the wordmark points at {cell.hyperlink.target!r}, not the platform"
    )


def test_the_cover_links_all_use_the_canonical_destination(cover):
    """One place defines where each destination is, so the cover cannot drift.

    The platform URL and the author's own site are both held in single constants
    precisely so a regenerated workbook cannot point at a hand-typed or stale
    destination. The author signature is a separate, deliberate destination and
    is not expected to match the platform's.
    """
    from backend.export.excel.links import AUTHOR_URL

    linked = {
        c.coordinate: c.hyperlink.target
        for row in cover.iter_rows()
        for c in row
        if c.hyperlink is not None and c.hyperlink.target
    }

    assert linked, "the cover carries no links at all"

    known = {VALENCE_URL.rstrip("/"), AUTHOR_URL.rstrip("/")}
    for coord, target in linked.items():
        assert target.rstrip("/") in known, (
            f"{coord} points at {target!r}, which is neither the canonical "
            f"platform URL nor the author's site"
        )

    # The wordmark is the single place that navigates to the platform; the
    # strapline beneath it is plain text, so the cover has one obvious target
    # rather than two competing ones. The signature goes to the author.
    assert linked.get("B2", "").rstrip("/") == VALENCE_URL.rstrip("/")
    assert linked.get("B20", "").rstrip("/") == AUTHOR_URL.rstrip("/")


def test_the_strapline_is_not_a_second_link(cover):
    """One target on the cover, not two.

    Two adjacent cells linking to the same destination makes the reader choose
    between them and neither looks like the intended one.
    """
    assert cover["B3"].hyperlink is None, (
        "the strapline under the wordmark is hyperlinked as well, so the cover "
        "offers two links to the same place"
    )
