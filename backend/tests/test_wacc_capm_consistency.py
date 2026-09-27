"""WACC must be reproducible from its own published CAPM components.

The engine recomputes WACC on every request against live market data, but
`wacc.cost_of_equity` also exists as a model_generated assumption captured when
the model was cached. Honouring that snapshot froze the discount rate: a US
model cached at a 4.64% risk-free rate kept discounting at 13.17% cost of equity
long after the live rate moved, and the exported workbook — which recomputes
CAPM from today's rf/beta/ERP — showed a different WACC than the API with no
visible cause.

These tests pin: CAPM is authoritative, only an analyst override may pin Ke,
and the Excel WACC chain reproduces the API number from its own cells.
"""

import pytest
from openpyxl import Workbook

from backend.export.excel.render_fcst import WACC_ROWS
from backend.export.excel.render_val import render_wacc_tab
from backend.models.spec.assumptions import AssumptionObject
from backend.models.spec.historicals import Historicals
from backend.models.spec.metadata import ModelMetadata
from backend.models.spec.model_specification import ModelSpecification
from backend.models.spec.valuation import ValuationOutput, WACCBreakdown
from backend.valuation.wacc import compute_wacc


@pytest.fixture(autouse=True)
def fixed_market_data(monkeypatch):
    """Pin the provider so the test never depends on a live quote."""
    from backend.data.providers import market_data as md

    def _pt(value, note="test"):
        return md.MarketDataPoint(value=value, source="test", provenance_note=note)

    class _Stub:
        risk_free_rate = _pt(5.0)
        beta = _pt(1.5)
        equity_risk_premium = _pt(4.5)
        price = _pt(200.0)
        shares_outstanding = _pt(1000.0)

    monkeypatch.setattr("backend.valuation.wacc.get_company_market_data", lambda *a, **k: _Stub())


def _assumptions(ke_value: float, ke_type: str = "model_generated") -> list[AssumptionObject]:
    return [
        AssumptionObject(
            driver_key="wacc.cost_of_equity",
            value=ke_value,
            period="all",
            scenario="base",
            type=ke_type,
            source="test",
        )
    ]


def test_model_generated_ke_snapshot_does_not_freeze_the_discount_rate():
    """A stale cached Ke must not override today's CAPM."""
    w = compute_wacc(
        assumptions=_assumptions(13.17),
        scenario="base",
        current_share_price=200.0,
        debt_cr=0.0,
    )
    # Provider-sourced beta is Blume-adjusted: 0.67*1.5 + 0.33 = 1.335.
    # CAPM: 5.0 + 1.335*4.5 = 11.0075
    assert w.beta == pytest.approx(1.335, abs=0.001)
    assert w.cost_of_equity == pytest.approx(5.0 + 1.335 * 4.5, abs=0.01)
    assert w.wacc == pytest.approx(5.0 + 1.335 * 4.5, abs=0.01)
    assert "= Rf + Beta x ERP" in w.source_notes


def test_analyst_override_still_pins_cost_of_equity():
    w = compute_wacc(
        assumptions=_assumptions(14.0, ke_type="user_override"),
        scenario="base",
        current_share_price=200.0,
        debt_cr=0.0,
    )
    assert w.cost_of_equity == pytest.approx(14.0, abs=0.01)
    assert w.wacc == pytest.approx(14.0, abs=0.01)
    assert "analyst override" in w.source_notes


def _spec_with_wacc(ke: float, rfr=5.0, beta=1.5, erp=4.5) -> ModelSpecification:
    spec = ModelSpecification(
        metadata=ModelMetadata(
            company_id="wacc_us",
            ticker="WCT",
            name="WACC Testco",
            market="us",
            currency="USD",
            units="millions",
            fiscal_year_end="December 31",
            shares_outstanding=1000.0,
        ),
        historicals=Historicals(periods=["FY26"], line_items=[]),
    )
    capm_ke = rfr + beta * erp
    wacc_b = WACCBreakdown(
        risk_free_rate=rfr,
        beta=beta,
        equity_risk_premium=erp,
        cost_of_equity=ke,
        pre_tax_cost_of_debt=7.5,
        tax_rate=21.0,
        cost_of_debt=7.5 * (1 - 0.21),
        equity_weight=0.99,
        debt_weight=0.01,
        wacc=(0.99 * ke + 0.01 * 7.5 * (1 - 0.21)),
        source_notes="test",
    )
    spec.valuation = [
        ValuationOutput(scenario=sc, wacc=wacc_b, dcf_bridge=__import__(
            "backend.models.spec.valuation", fromlist=["DCFBridge"]
        ).DCFBridge(enterprise_value=1.0, equity_value=1.0, shares_outstanding=1000.0, implied_share_price=1.0))
        for sc in ("base", "bull", "bear")
    ]
    assert capm_ke  # documents the CAPM identity the workbook must reproduce
    return spec


def _sheet_wacc(spec) -> float:
    """Recompute WACC from the workbook's own cells, as Excel would.

    Cells are addressed through the published WACC row map rather than by number.
    A hardcoded row made this test a tripwire for the wrong thing: when a line was
    inserted above it, the test read the wrong cells and failed for a reason that
    had nothing to do with what it claimed to check.
    """
    ws = render_wacc_tab(Workbook(), spec)

    def cell(line: str):
        return ws.cell(row=WACC_ROWS[line], column=3).value

    rfr, beta, erp = cell("risk_free_rate"), cell("beta_used"), cell("equity_risk_premium")
    kd_pre, tax = cell("pre_tax_cost_of_debt"), cell("tax_rate")
    we, wd = cell("equity_weight"), cell("debt_weight")
    ke_cell = cell("cost_of_equity")
    ke = float(ke_cell) if not str(ke_cell).startswith("=") else rfr + beta * erp
    kd = kd_pre * (1 - tax)
    return (we * ke + wd * kd) * 100.0


def test_excel_wacc_reproduces_the_api_number():
    for ke in (11.75, 14.0):
        spec = _spec_with_wacc(ke)
        api_wacc = spec.get_valuation("base").wacc.wacc
        assert _sheet_wacc(spec) == pytest.approx(api_wacc, abs=0.01), (
            f"workbook WACC diverges from API WACC for cost of equity {ke}"
        )


def test_excel_ke_cell_is_an_input_only_when_overridden():
    capm = render_wacc_tab(Workbook(), _spec_with_wacc(11.75))
    ke_cell = capm.cell(row=WACC_ROWS["cost_of_equity"], column=3).value
    assert str(ke_cell).startswith("="), "CAPM cost of equity must stay a live formula"

    overridden = render_wacc_tab(Workbook(), _spec_with_wacc(14.0))
    assert overridden.cell(row=WACC_ROWS["cost_of_equity"], column=3).value == pytest.approx(0.14)


def test_the_capm_formula_names_the_beta_that_was_actually_used():
    """The cost-of-equity formula must reference the adjusted beta, not the raw one.

    The two now sit on separate rows, and only one of them belongs in CAPM. A
    formula pointing at the other would compute a cost of equity the model never
    discounted at, while still reproducing CAPM's shape.
    """
    ws = render_wacc_tab(Workbook(), _spec_with_wacc(11.75))
    formula = str(ws.cell(row=WACC_ROWS["cost_of_equity"], column=3).value)

    assert f"C{WACC_ROWS['beta_used']}" in formula, (
        f"the CAPM formula does not reference the beta that was used: {formula}"
    )
    assert f"C{WACC_ROWS['beta_raw']}" not in formula, (
        f"the CAPM formula references the published beta rather than the "
        f"adjusted one: {formula}"
    )


def test_an_overridden_cost_of_equity_says_so_on_the_page():
    """The override must be visible in the provenance column, not just in the cell.

    A pinned cost of equity is an assumption, and an assumption that the page does
    not label reads as a result.
    """
    ws = render_wacc_tab(Workbook(), _spec_with_wacc(14.0))
    note = ws.cell(row=WACC_ROWS["cost_of_equity"], column=4).value

    assert "Analyst-set" in str(note), (
        f"the provenance column does not state that the cost of equity was "
        f"set by hand: {note!r}"
    )
